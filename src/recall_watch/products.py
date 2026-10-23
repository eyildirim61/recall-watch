"""Load the product list a business wants watched (CSV, or XLSX with openpyxl)."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from .gtin import item_from_case, normalize

ALIASES: Dict[str, Sequence[str]] = {
    "name": ("name", "product name", "product", "description", "item", "urun", "urun adi", "stok adi"),
    "brand": ("brand", "marka", "manufacturer", "uretici"),
    "sku": ("sku", "internal reference", "default_code", "item code", "product code", "stok kodu", "urun kodu"),
    "barcode": ("barcode", "upc", "ean", "gtin", "unit barcode", "barkod", "birim barkod"),
    "box_barcode": ("box barcode", "case barcode", "carton barcode", "itf14", "gtin14", "koli barkodu", "koli barkod"),
}

_TR = str.maketrans({"ı": "i", "ç": "c", "ş": "s", "ğ": "g", "ü": "u", "ö": "o"})


class ProductError(Exception):
    pass


def norm_text(text: str) -> str:
    text = unicodedata.normalize("NFKD", (text or "").casefold())
    text = "".join(c for c in text if not unicodedata.combining(c)).translate(_TR)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return text.strip()


def _header_key(h: str) -> str:
    return re.sub(r"\s*\([^)]*\)\s*$", "", norm_text(h).replace("_", " ")).strip()


@dataclass
class Product:
    row: int
    name: str = ""
    brand: str = ""
    sku: str = ""
    gtins: List[str] = field(default_factory=list)  # normalised GTIN-14, unit and box, plus derived items

    @property
    def label(self) -> str:
        parts = [p for p in (self.sku, self.name) if p]
        return " · ".join(parts) or f"row {self.row}"


def brand_from_name(name: str) -> str:
    """Many ERPs write "Brand, product, size": take the part before the first comma."""
    head = name.split(",", 1)[0].strip()
    return head if 0 < len(head) <= 40 and head != name.strip() else ""


def _read_rows(path: Path):
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise ProductError("Reading .xlsx needs openpyxl: pip install openpyxl (or save the list as CSV)") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [["" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
                     for v in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
        return rows
    data = path.read_bytes()
    for enc in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def load_products(path, brand_from_names: bool = True) -> List[Product]:
    path = Path(path)
    if not path.exists():
        raise ProductError(f"Product list not found: {path}")
    rows = _read_rows(path)
    if not rows:
        raise ProductError(f"{path.name} is empty")
    headers = [_header_key(h) for h in rows[0]]
    cols: Dict[str, int] = {}
    for key, aliases in ALIASES.items():
        for i, h in enumerate(headers):
            if h in aliases and i not in cols.values():
                cols[key] = i
                break
    if "name" not in cols and "barcode" not in cols and "brand" not in cols:
        raise ProductError(f"{path.name}: need at least a name, brand or barcode column. Found: {', '.join(rows[0])}")

    def get(r, key):
        i = cols.get(key)
        return r[i].strip() if i is not None and i < len(r) else ""

    products = []
    for n, r in enumerate(rows[1:], start=2):
        if not any(c.strip() for c in r):
            continue
        name, brand = get(r, "name"), get(r, "brand")
        name = re.sub(r"^\[[^\]]*\]\s*", "", name)  # Odoo's "[SKU] Name" display form
        if not brand and brand_from_names:
            brand = brand_from_name(name)
        gtins: List[str] = []
        for key in ("barcode", "box_barcode"):
            code = normalize(get(r, key).split("(")[0])
            if code:
                gtins.append(code)
                item = item_from_case(code)
                if key == "box_barcode" and item:
                    gtins.append(item)
        products.append(Product(row=n, name=name, brand=brand, sku=get(r, "sku"), gtins=list(dict.fromkeys(gtins))))
    return products
